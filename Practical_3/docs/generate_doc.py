import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = docx.Document()

# Set standard margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def add_title(text, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    return p

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(15, 23, 42)
    return h

def add_p(text, bold_prefix='', space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(12)
        r_b.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p

# --- COVER PAGE ---
for _ in range(2): 
    doc.add_paragraph()
add_title('MIT ACADEMY OF ENGINEERING', size=16, bold=True, space_after=4)
add_title('ALANDI, PUNE', size=13, bold=True, space_after=24)
add_title('WEB TECHNOLOGY', size=14, bold=True, space_after=6)
add_title('PRACTICAL NO. 3', size=16, bold=True, space_after=6)
add_title('FRONT-END JAVASCRIPT FRAMEWORKS: REACTJS', size=14, bold=True, space_after=36)

add_title('Project:', size=12, bold=False, space_after=6)
add_title('JOBAPPLY', size=18, bold=True, space_after=4)
add_title('JOB APPLICATION PORTAL', size=13, bold=True, space_after=48)

# Table for Student Meta
table = doc.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ('Student Name:', 'Pankaj Shinde'),
    ('PRN:', '202401100005'),
    ('Branch:', 'Software Engineering'),
    ('Division:', 'A')
]
for idx, (label, val) in enumerate(meta_data):
    row = table.rows[idx]
    row.cells[0].paragraphs[0].add_run(label).font.bold = True
    row.cells[1].paragraphs[0].add_run(val)
    for cell in row.cells:
        cell.paragraphs[0].runs[0].font.name = 'Times New Roman'
        cell.paragraphs[0].runs[0].font.size = Pt(11)

for _ in range(4): 
    doc.add_paragraph()
add_p('https://github.com/Pankaj-Shinde-07/JobApply-Practical-3', bold_prefix='GitHub Repository : ')
add_p('https://pankaj-shinde-07.github.io/JobApply-Practical-3/', bold_prefix='Deployment Link : ')

doc.add_page_break()

# --- BODY ---
add_heading('1. AIM', level=1)
add_p('To design, develop and deploy a ReactJS single-page web application named JobApply — Job Application Portal, demonstrating React components, state management using hooks, event handling, client-side form validation, error and success messages, interactive application management, and browser-side data persistence.')

add_heading('2. PROBLEM STATEMENT', level=1)
add_p('The practical requires an individual front-end web application developed using ReactJS. The application must demonstrate client-side validation and interactive functionality, use React concepts such as components, state and event handling, display suitable error messages, provide success or confirmation feedback, offer a user-friendly interface, and be tested with valid and invalid inputs. The completed application must also be deployed and submitted through a public link.')
add_p('JobApply addresses these requirements through a job application form. A user can enter personal and job-related information, validate details before submission, receive field-level errors when input is invalid, and receive a success message after a valid application is added. Submitted applications are displayed in an interactive list, can be deleted, and are saved in browser localStorage so they remain available after a refresh in the same browser.')

add_heading('3. OBJECTIVES', level=1)
objectives = [
    'Develop an individual ReactJS application using a component-based structure.',
    'Use React state (useState) to manage form values, validation errors, notifications and application records.',
    'Implement controlled form inputs and event handlers for user interaction.',
    'Validate name, email, Indian mobile number, job role, education, experience and optional resume URL.',
    'Display field-level validation errors and overall success or warning messages.',
    'Allow users to submit valid applications and view the application list dynamically.',
    'Provide a delete action for removing an application from the list.',
    'Persist application records in browser localStorage using JSON conversion.',
    'Create a responsive, user-friendly interface for desktop, tablet and mobile screens.',
    'Build and deploy the production application using Vite and GitHub Pages.'
]
for obj in objectives:
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(obj)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

add_heading('4. INTRODUCTION', level=1)
add_p('ReactJS is a JavaScript library used to build interactive user interfaces. It organizes an interface into reusable components and updates the rendered page when application state changes. React state and event handlers allow a web application to respond to user input without manually rebuilding the entire page.')
add_p('JobApply — Job Application Portal is a single-page front-end application created for Web Technology Practical 3. It provides a form for entering applicant information, selecting a job role and education level, entering experience, and optionally supplying a resume URL. The application validates the entered data before accepting it and presents feedback when the form is incomplete or invalid.')
add_p('The application uses React useState for state management and useEffect for synchronizing application records with localStorage. Vite is used as the development and production build tool. The project is hosted in a public GitHub repository and deployed to GitHub Pages.')

add_heading('5. TECHNOLOGIES USED', level=1)
tech_table = doc.add_table(rows=1, cols=2)
tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr_cells = tech_table.rows[0].cells
hdr_cells[0].paragraphs[0].add_run('Technology').font.bold = True
hdr_cells[1].paragraphs[0].add_run('Purpose').font.bold = True
for c in hdr_cells:
    c.paragraphs[0].runs[0].font.name = 'Times New Roman'

technologies = [
    ('ReactJS', 'Component-based user interface and declarative virtual DOM rendering.'),
    ('JavaScript (ES6+)', 'Validation logic, event callbacks, object and array operations, regex.'),
    ('useState', 'Manages form data, errors, notifications and application records.'),
    ('useEffect', 'Synchronizes saved application data with browser localStorage.'),
    ('HTML5 / JSX', 'Defines the interface structure using JSX syntax.'),
    ('CSS3', 'Custom visual styling, layout, colours, spacing, states and responsive refinements.'),
    ('Vite', 'Development server and production build tool.'),
    ('npm', 'Installs project dependencies and runs build/lint scripts.'),
    ('localStorage / JSON', 'Browser-side persistence of application records.'),
    ('VS Code / Web Browser', 'Development, editing, testing and demonstration environment.'),
    ('Git / GitHub', 'Version control and source-code hosting.'),
    ('GitHub Actions / GitHub Pages', 'Automated production build and public static-site deployment.')
]
for tech, purpose in technologies:
    row_cells = tech_table.add_row().cells
    r0 = row_cells[0].paragraphs[0].add_run(tech)
    r1 = row_cells[1].paragraphs[0].add_run(purpose)
    r0.font.name = 'Times New Roman'
    r1.font.name = 'Times New Roman'
    r0.font.size = Pt(10.5)
    r1.font.size = Pt(10.5)

add_heading('6. CLIENT-SIDE VALIDATION', level=1)
val_table = doc.add_table(rows=1, cols=3)
val_table.alignment = WD_TABLE_ALIGNMENT.CENTER
v_hdr = val_table.rows[0].cells
v_hdr[0].paragraphs[0].add_run('Field').font.bold = True
v_hdr[1].paragraphs[0].add_run('Validation rule').font.bold = True
v_hdr[2].paragraphs[0].add_run('Example of invalid input').font.bold = True

v_rules = [
    ('Full name', 'Trimmed name must contain at least 3 characters.', 'Blank or short name (e.g. "Ab").'),
    ('Email', 'Must match valid email pattern user@domain.tld.', 'name@ or missing top-level domain.'),
    ('Mobile number', '10-digit Indian mobile pattern beginning with 6-9.', '12345, letters, or starts with 0-5.'),
    ('Job role', 'A job role must be selected from list.', 'No selection ("Select a role").'),
    ('Education', 'An education value must be selected.', 'No selection ("Select education").'),
    ('Experience', 'Must be a finite number between 0 and 50.', '-1, 51, or non-numeric text.'),
    ('Resume URL', 'Optional; if supplied, must be valid HTTP or HTTPS URL.', 'ftp://link or invalid string.')
]
for field, rule, ex in v_rules:
    rc = val_table.add_row().cells
    rc[0].paragraphs[0].add_run(field).font.bold = True
    rc[1].paragraphs[0].add_run(rule)
    rc[2].paragraphs[0].add_run(ex)
    for c in rc:
        c.paragraphs[0].runs[0].font.name = 'Times New Roman'
        c.paragraphs[0].runs[0].font.size = Pt(10.5)

add_heading('7. TESTING AND VALIDATION (T1 to T14)', level=1)
test_table = doc.add_table(rows=1, cols=4)
test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
t_hdr = test_table.rows[0].cells
t_hdr[0].paragraphs[0].add_run('Test').font.bold = True
t_hdr[1].paragraphs[0].add_run('Input / Action').font.bold = True
t_hdr[2].paragraphs[0].add_run('Expected Result').font.bold = True
t_hdr[3].paragraphs[0].add_run('Observed Result').font.bold = True

tests = [
    ('T1', 'Submit all fields empty.', 'Required-field messages appear; no record is added.', 'PASS (Verified)'),
    ('T2', 'Enter a very short name.', 'Name validation message appears.', 'PASS (Verified)'),
    ('T3', 'Enter an invalid email address.', 'Email validation message appears.', 'PASS (Verified)'),
    ('T4', 'Enter an invalid mobile number.', 'Mobile validation message appears.', 'PASS (Verified)'),
    ('T5', 'Leave job role or education unselected.', 'Selection validation appears.', 'PASS (Verified)'),
    ('T6', 'Enter experience below 0 or above 50.', 'Experience validation appears.', 'PASS (Verified)'),
    ('T7', 'Enter valid details and submit.', 'Application is added and success message appears.', 'PASS (Verified)'),
    ('T8', 'Leave resume URL blank.', 'Allowed if other fields are valid.', 'PASS (Verified)'),
    ('T9', 'Enter an invalid resume URL.', 'Resume URL validation appears.', 'PASS (Verified)'),
    ('T10', 'Refresh after a successful submission.', 'Saved application remains in the same browser.', 'PASS (Verified)'),
    ('T11', 'Delete an application.', 'Record disappears and count updates.', 'PASS (Verified)'),
    ('T12', 'Clear the form.', 'Form fields reset.', 'PASS (Verified)'),
    ('T13', 'Delete all records.', 'Empty-state message appears.', 'PASS (Verified)'),
    ('T14', 'Open at mobile viewport width.', 'Layout remains readable without horizontal overflow.', 'PASS (Verified)')
]
for tid, action, exp, obs in tests:
    rc = test_table.add_row().cells
    rc[0].paragraphs[0].add_run(tid).font.bold = True
    rc[1].paragraphs[0].add_run(action)
    rc[2].paragraphs[0].add_run(exp)
    rc[3].paragraphs[0].add_run(obs)
    for c in rc:
        c.paragraphs[0].runs[0].font.name = 'Times New Roman'
        c.paragraphs[0].runs[0].font.size = Pt(10)

doc.add_page_break()
add_heading('8. SCREENSHOTS', level=1)
img_folder = os.path.join(os.path.dirname(__file__), 'screenshots')
screenshots = [
    ('Figure_1_Landing_Page.png', 'Figure 1: JobApply desktop landing interface with application form and activity list.'),
    ('Figure_2_Validation_Errors.png', 'Figure 2: Real-time client-side validation error messages and highlights for incomplete input.'),
    ('Figure_3_Submission_Success.png', 'Figure 3: Success state, toast notification, and dynamically rendered application card.'),
    ('Figure_4_Test_Checklist_Modal.png', 'Figure 4: Interactive Practical 3 Test Suite modal (T1 to T14).'),
    ('Figure_5_JSON_State_Modal.png', 'Figure 5: LocalStorage persistent state inspected in formatted JSON format.')
]
for fname, caption in screenshots:
    fpath = os.path.join(img_folder, fname)
    if os.path.exists(fpath):
        doc.add_paragraph().paragraph_format.space_before = Pt(8)
        doc.add_picture(fpath, width=Inches(6.2))
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(14)
        cpr = cp.add_run(caption)
        cpr.font.name = 'Times New Roman'
        cpr.font.size = Pt(10)
        cpr.font.italic = True

add_heading('9. CONCLUSION', level=1)
add_p('The JobApply project demonstrates the development of a ReactJS front-end application using JSX, React state, event handling, client-side validation and dynamic rendering. The form validates applicant details before accepting a record, displays useful error feedback for invalid input, and confirms successful submission. Application records are displayed interactively and persisted in localStorage for the same browser.')
add_p('The project also demonstrates a Vite production build, source-code management through Git and GitHub, and automated deployment through GitHub Actions to GitHub Pages.')

out_path = 'd:/TY subjects/WT Lab/Practical 3.docx'
out_path_prn = 'd:/TY subjects/WT Lab/Practical 3-202401100005.docx'
doc.save(out_path)
doc.save(out_path_prn)
print('Successfully saved report to', out_path, 'and', out_path_prn)
