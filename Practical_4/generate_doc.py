import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = docx.Document()

# 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def add_title(text, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    return p

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(15, 23, 42)
    return h

def add_p(text, bold_prefix='', space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(12)
        r_b.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p

# ─── COVER PAGE ───────────────────────────────────────────────────────────────
for _ in range(2): 
    doc.add_paragraph()

add_title('MIT ACADEMY OF ENGINEERING', size=16, bold=True, space_after=4)
add_title('ALANDI, PUNE', size=13, bold=True, space_after=24)
add_title('DEPARTMENT OF SOFTWARE ENGINEERING', size=13, bold=True, space_after=6)
add_title('WEB TECHNOLOGY LABORATORY (WT LAB)', size=14, bold=True, space_after=6)
add_title('PRACTICAL NO. 04', size=16, bold=True, space_after=6)
add_title('SERVER-SIDE DEVELOPMENT: NODEJS', size=14, bold=True, space_after=36)

add_title('Project:', size=12, bold=False, space_after=6)
add_title('NODEREGISTER', size=18, bold=True, space_after=4)
add_title('STUDENT ENROLLMENT PORTAL WITH PURE NODEJS SERVER-SIDE VALIDATION', size=12, bold=True, space_after=40)

# Student Meta Table
table = doc.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ('Student Name:', 'Pankaj Shinde'),
    ('PRN:', '202401100005'),
    ('Branch:', 'Software Engineering'),
    ('Division:', 'A')
]
for idx, (label, val) in enumerate(meta_data):
    row = table.rows[idx]
    row.cells[0].paragraphs[0].add_run(label).font.bold = True
    row.cells[1].paragraphs[0].add_run(val)
    for cell in row.cells:
        cell.paragraphs[0].runs[0].font.name = 'Times New Roman'
        cell.paragraphs[0].runs[0].font.size = Pt(11)

for _ in range(4): 
    doc.add_paragraph()

add_p('Pure NodeJS (http, url, fs, path, querystring core modules - No Express)', bold_prefix='Server Technology: ')
add_p('http://localhost:4000', bold_prefix='Local Execution URL: ')

doc.add_page_break()

# ─── BODY ─────────────────────────────────────────────────────────────────────

add_heading('1. AIM', level=1)
add_p('To design, develop and demonstrate an individual web application named NodeRegister demonstrating server-side validation using pure NodeJS (without third-party frameworks), handling GET and POST HTTP streaming requests, request/response headers, status codes, modular routing, and persistent JSON storage.')

add_heading('2. PROBLEM STATEMENT', level=1)
add_p('Client-side validation alone is insufficient for secure web applications because malicious clients or API tools (such as Postman or curl) can easily bypass browser checks. Academic enrollment systems require rigorous server-side validation to ensure that student PRN numbers, institutional email addresses, mobile numbers, department assignments, and CGPA scores strictly conform to organizational criteria before records are accepted.')
add_p('NodeRegister solves this challenge using pure Node.js. An HTTP server created via http.createServer listens for incoming chunked data streams, parses URL-encoded and JSON bodies, executes modular validation rules in modules/validator.js, rejects malformed payloads with HTTP 400 Bad Request and detailed error messages, and commits valid enrollments into data/students.json with HTTP 201 Created.')

add_heading('3. OBJECTIVES', level=1)
objs = [
    'Develop an individual web application using pure NodeJS core modules.',
    'Implement comprehensive server-side validation for all incoming student input fields.',
    'Handle HTTP GET and POST requests appropriately using native chunked request streaming.',
    'Return proper success (HTTP 201) and error (HTTP 400, 404) status codes and descriptive feedback.',
    'Utilize NodeJS modular concepts by segregating code into validator, storage, and router modules.',
    'Test and verify the application using different valid and invalid test inputs.'
]
for o in objs:
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(o)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)

add_heading('4. TECHNOLOGIES USED', level=1)
tech_table = doc.add_table(rows=1, cols=2)
tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tech_table.rows[0].cells
hdr[0].paragraphs[0].add_run('Node.js Module / Tool').font.bold = True
hdr[1].paragraphs[0].add_run('Purpose & Functionality').font.bold = True
for c in hdr:
    c.paragraphs[0].runs[0].font.name = 'Times New Roman'

technologies = [
    ('http (Node.js Core Module)', 'Creates the HTTP web server (http.createServer) and handles incoming requests and outbound responses.'),
    ('url (Node.js Core Module)', 'Parses requested URL paths, query parameters, and endpoints for request routing.'),
    ('fs (Node.js Core Module)', 'Performs file reading and writing to persistently store registered student records in students.json.'),
    ('path (Node.js Core Module)', 'Constructs normalized file paths across different operating systems safely.'),
    ('querystring (Node.js Core)', 'Parses URL-encoded form submissions from POST request body streams.'),
    ('HTML5, CSS3 & JavaScript', 'Provides the responsive front-end user interface, demo quick-fill buttons, and dynamic roster table.')
]
for t, p in technologies:
    row = tech_table.add_row().cells
    r0 = row[0].paragraphs[0].add_run(t)
    r1 = row[1].paragraphs[0].add_run(p)
    r0.font.name = 'Times New Roman'
    r1.font.name = 'Times New Roman'
    r0.font.size = Pt(10)
    r1.font.size = Pt(10)

add_heading('5. SERVER-SIDE VALIDATION MATRIX', level=1)
val_table = doc.add_table(rows=1, cols=4)
val_table.alignment = WD_TABLE_ALIGNMENT.CENTER
v_hdr = val_table.rows[0].cells
v_hdr[0].paragraphs[0].add_run('Input Field').font.bold = True
v_hdr[1].paragraphs[0].add_run('Server-Side Rule').font.bold = True
v_hdr[2].paragraphs[0].add_run('Invalid Test Input').font.bold = True
v_hdr[3].paragraphs[0].add_run('Server Error Message Returned (HTTP 400)').font.bold = True
for c in v_hdr:
    c.paragraphs[0].runs[0].font.name = 'Times New Roman'

val_rules = [
    ('Student Name', 'Length >= 3 chars; letters, spaces and dots only', '"Al" or blank', 'Full name must contain at least 3 characters.'),
    ('12-Digit PRN', 'Exact 12-digit numeric regex: /^\\d{12}$/', '"12345" (5 digits)', 'PRN must consist of exactly 12 numeric digits.'),
    ('Email Address', 'Standard email pattern /^[^\\s@]+@[^\\s@]+\\.[^\\s@]+$/', '"invalid-email"', 'Please provide a valid email format.'),
    ('Mobile Number', '10-digit Indian phone starting 6-9: /^[6-9]\\d{9}$/', '"123456"', 'Enter a valid 10-digit Indian mobile number beginning with 6-9.'),
    ('Department', 'Non-empty; must match approved college departments', 'Blank or unknown', 'Academic department selection is required.'),
    ('Semester', 'Non-empty academic semester', 'Blank ("")', 'Academic semester selection is required.'),
    ('Elective Course', 'Must match approved course offering list', 'Blank ("")', 'Elective course selection is required.'),
    ('Cumulative CGPA', 'Float between 0.00 and 10.00', '"15.5" or negative', 'CGPA must be a valid numeric score between 0.00 and 10.00.')
]
for f, r, ex, em in val_rules:
    row = val_table.add_row().cells
    row[0].paragraphs[0].add_run(f).font.bold = True
    row[1].paragraphs[0].add_run(r)
    row[2].paragraphs[0].add_run(ex)
    row[3].paragraphs[0].add_run(em)
    for c in row:
        c.paragraphs[0].runs[0].font.name = 'Times New Roman'
        c.paragraphs[0].runs[0].font.size = Pt(9.5)

doc.add_page_break()

add_heading('6. TESTING AND VALIDATION (T1 to T10)', level=1)
test_table = doc.add_table(rows=1, cols=4)
test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
t_hdr = test_table.rows[0].cells
t_hdr[0].paragraphs[0].add_run('Test').font.bold = True
t_hdr[1].paragraphs[0].add_run('Input / Action').font.bold = True
t_hdr[2].paragraphs[0].add_run('Expected Node.js Server Response').font.bold = True
t_hdr[3].paragraphs[0].add_run('Observed Result').font.bold = True
for c in t_hdr:
    c.paragraphs[0].runs[0].font.name = 'Times New Roman'

tests = [
    ('T1', 'Submit completely empty form payload', 'HTTP 400 Bad Request; returns errors for all 8 required fields', 'PASS (Verified)'),
    ('T2', 'Enter short name ("Al")', 'HTTP 400; "Full name must contain at least 3 characters."', 'PASS (Verified)'),
    ('T3', 'Enter 5-digit PRN ("12345")', 'HTTP 400; PRN 12-digit regex error displayed', 'PASS (Verified)'),
    ('T4', 'Enter invalid email ("student@domain")', 'HTTP 400; Invalid email regex format error displayed', 'PASS (Verified)'),
    ('T5', 'Enter non-Indian mobile ("123456")', 'HTTP 400; Indian 10-digit mobile validation message shown', 'PASS (Verified)'),
    ('T6', 'Leave department or course unselected', 'HTTP 400; Required dropdown selection error displayed', 'PASS (Verified)'),
    ('T7', 'Enter out-of-range CGPA ("15.5")', 'HTTP 400; CGPA numeric range 0.00-10.00 error displayed', 'PASS (Verified)'),
    ('T8', 'Click "Fill Valid Student Demo" & Submit', 'HTTP 201 Created; record appended; verified receipt card issued', 'PASS (Verified)'),
    ('T9', 'Refresh browser page', 'Saved student records remain intact from students.json', 'PASS (Verified)'),
    ('T10', 'Click Delete button on student record', 'HTTP 200 OK; record removed from table and JSON file', 'PASS (Verified)')
]
for tid, act, exp, obs in tests:
    row = test_table.add_row().cells
    row[0].paragraphs[0].add_run(tid).font.bold = True
    row[1].paragraphs[0].add_run(act)
    row[2].paragraphs[0].add_run(exp)
    row[3].paragraphs[0].add_run(obs)
    for c in row:
        c.paragraphs[0].runs[0].font.name = 'Times New Roman'
        c.paragraphs[0].runs[0].font.size = Pt(9.5)

add_heading('7. NODE.JS ARCHITECTURE & STREAM HANDLING', level=1)
add_p('In pure Node.js, request bodies arrive as asynchronous chunks of binary buffer data. The server buffers these chunks in memory before parsing:')
add_p('// Representative chunked stream handler:\nlet body = "";\nreq.on("data", chunk => { body += chunk.toString(); });\nreq.on("end", () => {\n  const data = JSON.parse(body);\n  const validation = validateStudent(data);\n  if (!validation.isValid) {\n    res.writeHead(400, { "Content-Type": "application/json" });\n    return res.end(JSON.stringify({ success: false, errors: validation.errors }));\n  }\n  // Persist record and return HTTP 201 Created\n});')

doc.add_page_break()

add_heading('8. SCREENSHOTS AND APPLICATION VERIFICATION', level=1)
img_folder = os.path.join(os.path.dirname(__file__), 'screenshots')
screenshots = [
    ('Figure_1_Landing_Page.png', 'Figure 1: NodeRegister enrollment interface with empty form controls and architecture summary.'),
    ('Figure_2_Validation_Errors.png', 'Figure 2: Pure Node.js server-side validation error messages and highlights on HTTP 400 rejection.'),
    ('Figure_3_Submission_Success.png', 'Figure 3: Confirmed student enrollment with HTTP 201 Created banner and verified registration receipt.'),
    ('Figure_4_Students_Roster.png', 'Figure 4: Registered Students Roster dynamically loaded from data/students.json on the server.')
]
for fname, caption in screenshots:
    fpath = os.path.join(img_folder, fname)
    if os.path.exists(fpath):
        doc.add_paragraph().paragraph_format.space_before = Pt(8)
        doc.add_picture(fpath, width=Inches(6.2))
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(14)
        cpr = cp.add_run(caption)
        cpr.font.name = 'Times New Roman'
        cpr.font.size = Pt(10)
        cpr.font.italic = True

add_heading('9. CONCLUSION', level=1)
add_p('Practical No. 04 successfully demonstrates the development of an individual server-side web application using pure NodeJS without third-party frameworks. The application demonstrates core Node.js concepts including http.createServer, chunked request streaming, request routing, header manipulation, HTTP status code management (200, 201, 400, 404), modular server-side validation, and persistent file-system storage in JSON format. All test cases were executed and verified.')

out_path = 'd:/TY subjects/WT Lab/Practical 4-202401100005.docx'
doc.save(out_path)
print('Successfully saved Practical 4 report to', out_path)
