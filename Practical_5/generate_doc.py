import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = docx.Document()

# 1-inch margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def add_title(text, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    return p

def add_heading(text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    for run in h.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(15, 23, 42)
    return h

def add_p(text, bold_prefix='', space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = 'Times New Roman'
        r_b.font.size = Pt(12)
        r_b.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p

# ─── COVER PAGE ───────────────────────────────────────────────────────────────
for _ in range(2): 
    doc.add_paragraph()

add_title('MIT ACADEMY OF ENGINEERING', size=16, bold=True, space_after=4)
add_title('ALANDI, PUNE', size=13, bold=True, space_after=24)
add_title('DEPARTMENT OF SOFTWARE ENGINEERING', size=13, bold=True, space_after=6)
add_title('WEB TECHNOLOGY LABORATORY (WT LAB)', size=14, bold=True, space_after=6)
add_title('PRACTICAL NO. 05', size=16, bold=True, space_after=6)
add_title('WEB APPLICATION FRAMEWORK: EXPRESSJS', size=14, bold=True, space_after=36)

add_title('Project:', size=12, bold=False, space_after=6)
add_title('HEALTHCARE EXPRESS', size=18, bold=True, space_after=4)
add_title('PATIENT APPOINTMENT PORTAL WITH EXPRESSJS SERVER-SIDE VALIDATION', size=12, bold=True, space_after=40)

# Student Meta Table
table = doc.add_table(rows=4, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ('Student Name:', 'Pankaj Shinde'),
    ('PRN:', '202401100005'),
    ('Branch:', 'Software Engineering'),
    ('Division:', 'A')
]
for idx, (label, val) in enumerate(meta_data):
    row = table.rows[idx]
    row.cells[0].paragraphs[0].add_run(label).font.bold = True
    row.cells[1].paragraphs[0].add_run(val)
    for cell in row.cells:
        cell.paragraphs[0].runs[0].font.name = 'Times New Roman'
        cell.paragraphs[0].runs[0].font.size = Pt(11)

for _ in range(4): 
    doc.add_paragraph()

add_p('Express.js Framework (Express Routers, Custom Validation Middleware, REST APIs)', bold_prefix='Server Framework: ')
add_p('http://localhost:5000', bold_prefix='Local Execution URL: ')

doc.add_page_break()

# ─── BODY ─────────────────────────────────────────────────────────────────────

add_heading('1. AIM', level=1)
add_p('To design, develop and demonstrate an individual web application named HealthCare Express to demonstrate server-side validation using the ExpressJS framework, demonstrating modular routes, middleware request processing, request/response handling, HTTP status codes, and persistent data storage.')

add_heading('2. PROBLEM STATEMENT', level=1)
add_p('Clinical appointment booking systems require strict server-side validation to prevent scheduling conflicts, reject past booking dates, validate patient contact numbers and ages, and ensure that only clean, verified medical consult requests enter the clinical database. Relying solely on client-side validation is insecure as HTTP requests can be forged directly.')
add_p('HealthCare Express demonstrates how Express.js solves this through custom middleware. Before reaching the route handler, incoming requests pass through validateAppointmentMiddleware, which inspects patient details, validates date chronology (disallowing past dates), enforces Indian mobile regex and email patterns, and returns HTTP 400 Bad Request with field-specific errors upon validation failure, or passes sanitized data to the router to persist bookings in data/appointments.json.')

add_heading('3. OBJECTIVES', level=1)
objs = [
    'Develop an individual web application using ExpressJS framework.',
    'Implement robust server-side validation middleware for all user input fields.',
    'Create suitable modular Express routes to handle client requests (/api/appointments, /api/stats).',
    'Use ExpressJS for request and response handling with proper HTTP status codes (200, 201, 400, 404, 500).',
    'Display proper success and error feedback banners and field-specific validation messages.',
    'Test the application thoroughly using different valid and invalid test inputs.'
]
for o in objs:
    p = doc.add_paragraph(style='List Bullet')
    r = p.add_run(o)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11.5)

add_heading('4. TECHNOLOGIES USED', level=1)
tech_table = doc.add_table(rows=1, cols=2)
tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = tech_table.rows[0].cells
hdr[0].paragraphs[0].add_run('Technology / Tool').font.bold = True
hdr[1].paragraphs[0].add_run('Role & Purpose in Project').font.bold = True
for c in hdr:
    c.paragraphs[0].runs[0].font.name = 'Times New Roman'

technologies = [
    ('Express.js (v5)', 'Fast, unopinionated web framework providing routing, middleware pipelines, and HTTP utilities.'),
    ('Express Router', 'Modular, mountable route handlers separating appointment management and statistics.'),
    ('Express Middleware', 'Custom validation middleware intercepting requests before route execution.'),
    ('cors & body-parser', 'Middleware parsing incoming JSON and URL-encoded request bodies and enabling CORS.'),
    ('JSON File Storage (appointments.json)', 'Persistent file-backed database storing confirmed patient appointments.'),
    ('HTML5, CSS3 & JavaScript (ES6+)', 'Modern clinical user interface with responsive CSS grid and asynchronous fetch API.')
]
for t, p in technologies:
    row = tech_table.add_row().cells
    r0 = row[0].paragraphs[0].add_run(t)
    r1 = row[1].paragraphs[0].add_run(p)
    r0.font.name = 'Times New Roman'
    r1.font.name = 'Times New Roman'
    r0.font.size = Pt(10)
    r1.font.size = Pt(10)

add_heading('5. EXPRESS SERVER-SIDE VALIDATION MATRIX', level=1)
val_table = doc.add_table(rows=1, cols=4)
val_table.alignment = WD_TABLE_ALIGNMENT.CENTER
v_hdr = val_table.rows[0].cells
v_hdr[0].paragraphs[0].add_run('Field').font.bold = True
v_hdr[1].paragraphs[0].add_run('Express Middleware Validation Rule').font.bold = True
v_hdr[2].paragraphs[0].add_run('Invalid Test Example').font.bold = True
v_hdr[3].paragraphs[0].add_run('Error Message Displayed (HTTP 400)').font.bold = True
for c in v_hdr:
    c.paragraphs[0].runs[0].font.name = 'Times New Roman'

val_rules = [
    ('Patient Name', 'Length >= 3 chars; letters, spaces and periods only', '"P" or blank', 'Patient name must contain at least 3 characters.'),
    ('Patient Age', 'Integer between 1 and 120 years', '"150" or negative', 'Age must be between 1 and 120 years.'),
    ('Gender', 'Must be one of: Male, Female, Other', 'Blank or invalid', 'Gender selection is required.'),
    ('Email Address', 'Standard email pattern: /^[^\\s@]+@[^\\s@]+\\.[^\\s@]+$/', '"invalid-email"', 'Please provide a valid email address.'),
    ('Mobile Number', '10-digit Indian mobile regex: /^[6-9]\\d{9}$/', '"12345"', 'Enter a valid 10-digit Indian mobile number beginning with 6-9.'),
    ('Department', 'Must match approved clinical departments list', 'Blank ("")', 'Specialist medical department selection is required.'),
    ('Appointment Date', 'Must be today or a future date (no past dates!)', '"2020-01-01" (Past)', 'Appointment date cannot be in the past.'),
    ('Time Slot', 'Non-empty valid clinical time slot', 'Blank ("")', 'Preferred appointment time slot is required.'),
    ('Symptoms Detail', 'Minimum 10 characters descriptive text', '"Pain" (< 10 chars)', 'Please describe symptoms with at least 10 characters of detail.')
]
for f, r, ex, em in val_rules:
    row = val_table.add_row().cells
    row[0].paragraphs[0].add_run(f).font.bold = True
    row[1].paragraphs[0].add_run(r)
    row[2].paragraphs[0].add_run(ex)
    row[3].paragraphs[0].add_run(em)
    for c in row:
        c.paragraphs[0].runs[0].font.name = 'Times New Roman'
        c.paragraphs[0].runs[0].font.size = Pt(9.5)

doc.add_page_break()

add_heading('6. TESTING AND VALIDATION (T1 to T10)', level=1)
test_table = doc.add_table(rows=1, cols=4)
test_table.alignment = WD_TABLE_ALIGNMENT.CENTER
t_hdr = test_table.rows[0].cells
t_hdr[0].paragraphs[0].add_run('Test').font.bold = True
t_hdr[1].paragraphs[0].add_run('Input / Action').font.bold = True
t_hdr[2].paragraphs[0].add_run('Expected Express Server Response').font.bold = True
t_hdr[3].paragraphs[0].add_run('Observed Result').font.bold = True
for c in t_hdr:
    c.paragraphs[0].runs[0].font.name = 'Times New Roman'

tests = [
    ('T1', 'Submit all fields empty', 'HTTP 400 Bad Request; field errors returned for all required fields', 'PASS (Verified)'),
    ('T2', 'Enter short name ("P")', 'HTTP 400; "Patient name must contain at least 3 characters."', 'PASS (Verified)'),
    ('T3', 'Enter age > 120 ("150")', 'HTTP 400; Age range 1-120 error displayed', 'PASS (Verified)'),
    ('T4', 'Enter invalid email format', 'HTTP 400; Email regex validation error returned', 'PASS (Verified)'),
    ('T5', 'Enter non-Indian mobile ("12345")', 'HTTP 400; 10-digit Indian phone regex error displayed', 'PASS (Verified)'),
    ('T6', 'Select past date ("2020-01-01")', 'HTTP 400; "Appointment date cannot be in the past."', 'PASS (Verified)'),
    ('T7', 'Enter symptoms < 10 characters', 'HTTP 400; Minimum 10 characters requirement displayed', 'PASS (Verified)'),
    ('T8', 'Click "Fill Valid Patient Demo" & Submit', 'HTTP 201 Created; doctor assigned; receipt card displayed', 'PASS (Verified)'),
    ('T9', 'Change status dropdown in schedule', 'HTTP 200 OK via PATCH /api/appointments/:id/status', 'PASS (Verified)'),
    ('T10', 'Click Cancel on appointment', 'HTTP 200 OK via DELETE /api/appointments/:id; record removed', 'PASS (Verified)')
]
for tid, act, exp, obs in tests:
    row = test_table.add_row().cells
    row[0].paragraphs[0].add_run(tid).font.bold = True
    row[1].paragraphs[0].add_run(act)
    row[2].paragraphs[0].add_run(exp)
    row[3].paragraphs[0].add_run(obs)
    for c in row:
        c.paragraphs[0].runs[0].font.name = 'Times New Roman'
        c.paragraphs[0].runs[0].font.size = Pt(9.5)

add_heading('7. EXPRESS ROUTING & MIDDLEWARE ARCHITECTURE', level=1)
add_p('Express middleware functions have access to the request object (req), the response object (res), and the next middleware function in the application’s request-response cycle:')
add_p('// Representative validation middleware in Express:\nfunction validateAppointmentMiddleware(req, res, next) {\n  const { patientName, age, email, phone, date } = req.body;\n  const errors = {};\n  // Check fields and populate errors\n  if (Object.keys(errors).length > 0) {\n    return res.status(400).json({ success: false, errors });\n  }\n  req.validatedAppointment = req.body;\n  next(); // Pass control to next handler\n}')

doc.add_page_break()

add_heading('8. SCREENSHOTS AND APPLICATION VERIFICATION', level=1)
img_folder = os.path.join(os.path.dirname(__file__), 'screenshots')
screenshots = [
    ('Figure_1_Landing_Page.png', 'Figure 1: HealthCare Express clinical appointment booking landing page with doctor selection.'),
    ('Figure_2_Validation_Errors.png', 'Figure 2: Express server-side validation error messages and highlights on HTTP 400 rejection.'),
    ('Figure_3_Submission_Success.png', 'Figure 3: Confirmed medical consultation with HTTP 201 Created banner and verified booking receipt.'),
    ('Figure_4_Appointments_Schedule.png', 'Figure 4: Clinical Appointments Schedule table with doctor assignments and status pipeline dropdowns.')
]
for fname, caption in screenshots:
    fpath = os.path.join(img_folder, fname)
    if os.path.exists(fpath):
        doc.add_paragraph().paragraph_format.space_before = Pt(8)
        doc.add_picture(fpath, width=Inches(6.2))
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cp.paragraph_format.space_after = Pt(14)
        cpr = cp.add_run(caption)
        cpr.font.name = 'Times New Roman'
        cpr.font.size = Pt(10)
        cpr.font.italic = True

add_heading('9. CONCLUSION', level=1)
add_p('Practical No. 05 successfully demonstrates the development of an individual web application using the ExpressJS framework. The application implements modular routing with Express Router, custom server-side validation middleware, structured HTTP request/response handling with proper status codes (200, 201, 400, 404, 500), JSON persistence, and dynamic patient schedule management. All 10 test cases were executed and verified.')

out_path = 'd:/TY subjects/WT Lab/Practical 5-202401100005.docx'
doc.save(out_path)
print('Successfully saved Practical 5 report to', out_path)
